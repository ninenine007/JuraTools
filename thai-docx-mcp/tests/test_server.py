"""The tool surface: what a client sees, and what reaches the disk."""

import asyncio
import os
import sys
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from mcp.server.mcpserver.exceptions import ToolError  # noqa: E402

from thai_docx_mcp.server import server  # noqa: E402


def call(name, arguments):
    return asyncio.run(server.call_tool(name, arguments))


class Schemas(unittest.TestCase):
    def test_both_tools_are_registered_with_schemas(self):
        tools = {t.name: t for t in asyncio.run(server.list_tools())}
        self.assertEqual(set(tools), {"create_thai_document", "check_document_setup"})

        create = tools["create_thai_document"]
        self.assertIn("paragraphs", create.input_schema["properties"])
        self.assertIn("Block", create.input_schema["$defs"])
        # The model has to be told these are optional, or it invents values.
        self.assertEqual(create.input_schema.get("required"), ["paragraphs"])
        self.assertIn("location", create.output_schema["properties"])


class Generating(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.out = Path(self.tmp.name)
        os.environ["JURATOOLS_OUT_DIR"] = str(self.out)
        self.addCleanup(self.tmp.cleanup)
        self.addCleanup(os.environ.pop, "JURATOOLS_OUT_DIR", None)

    def test_a_document_is_written_with_the_title_as_a_heading(self):
        result = call("create_thai_document", {
            "title": "พินัยกรรม",
            "paragraphs": [{"text": "2.2. ทรัพย์มรดกฯ จำนวน 1,000.50 บาท"}],
        })
        location = Path(result.structured_content["location"])
        self.assertTrue(location.exists())
        self.assertEqual(location.name, "พินัยกรรม.docx")
        self.assertEqual(result.structured_content["paragraph_count"], 2)

        document = Document(str(location))
        heading, body = document.paragraphs
        self.assertEqual(heading.text, "พินัยกรรม")
        self.assertEqual(
            heading._p.find(qn("w:pPr")).find(qn("w:jc")).get(qn("w:val")), "center")
        self.assertIsNotNone(heading.runs[0]._element.rPr.find(qn("w:b")))
        self.assertEqual(
            body._p.find(qn("w:pPr")).find(qn("w:jc")).get(qn("w:val")), "thaiDistribute")

    def test_building_without_a_template_says_so(self):
        result = call("create_thai_document", {"paragraphs": [{"text": "ทดสอบ"}]})
        self.assertIsNone(result.structured_content["template_used"])
        self.assertTrue(any("template" in w.lower() for w in result.structured_content["warnings"]))

    def test_a_template_is_reported_and_leaves_no_warning(self):
        template_path = self.out / "ABC.docx"
        template = Document()
        template.add_paragraph("to be cleared")
        template.save(str(template_path))

        result = call("create_thai_document", {
            "paragraphs": [{"text": "ทดสอบ"}],
            "template": str(template_path),
            "file_name": "จากแม่แบบ",
        })
        self.assertEqual(result.structured_content["template_used"], str(template_path))
        self.assertEqual(result.structured_content["warnings"], [])
        self.assertEqual([p.text for p in Document(str(self.out / "จากแม่แบบ.docx")).paragraphs], ["ทดสอบ"])

    def test_a_missing_template_is_an_error_not_a_silent_fallback(self):
        """A template named but not there is a mistake worth stopping on — the
        alternative is a document quietly missing the firm's fonts."""
        with self.assertRaises(ToolError) as raised:
            call("create_thai_document", {
                "paragraphs": [{"text": "ทดสอบ"}],
                "template": str(self.out / "nope.docx"),
            })
        self.assertIn("nope.docx", str(raised.exception))

    def test_setup_reports_the_output_folder(self):
        result = call("check_document_setup", {})
        self.assertEqual(result.structured_content["output_directory"], str(self.out))
        self.assertFalse(result.structured_content["template_configured"])
        self.assertEqual(result.structured_content["default_font"], "Cordia New")


if __name__ == "__main__":
    unittest.main()
