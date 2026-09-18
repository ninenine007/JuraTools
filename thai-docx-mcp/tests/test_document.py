"""What the engine has to get right, checked in the XML itself.

A .docx that opens is not the same as a .docx a Thai lawyer can use, so these
assert on the markup — w:cs, thaiDistribute, the half-point sizes — rather than
on the file merely being produced.
"""

import sys
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from thai_docx_mcp.document import (  # noqa: E402
    Block,
    build_document,
    file_name_of,
    free_path,
    save_document,
)

THAI = "ข้อ 2.2. ผู้รับพินัยกรรมได้รับ 1,000.50 บาท ตามคดีหมายเลข 2569/2026"


def runs(paragraph):
    return [(r.text, r._element.rPr) for r in paragraph.runs]


class RunSplitting(unittest.TestCase):
    def setUp(self):
        document, _ = build_document([Block(text=THAI)])
        self.paragraph = document.paragraphs[0]

    def test_numbers_are_their_own_runs(self):
        """A trailing separator stays with the text beside it — "2.2." splits as
        the number "2.2" and a Thai run beginning with the full stop, which is
        how the prototype behaved and renders identically."""
        texts = [text for text, _ in runs(self.paragraph)]
        self.assertIn("2.2", texts)
        self.assertIn("1,000.50", texts)
        self.assertIn("2569/2026", texts)
        self.assertEqual(THAI, "".join(texts))

    def test_thai_runs_are_marked_complex_script(self):
        for text, properties in runs(self.paragraph):
            marked = properties.find(qn("w:cs")) is not None
            is_number = text.strip() and all(c in "0123456789.,/" for c in text)
            self.assertEqual(marked, not is_number, f"w:cs wrong on {text!r}")

    def test_every_run_carries_the_font_and_both_sizes(self):
        for text, properties in runs(self.paragraph):
            fonts = properties.find(qn("w:rFonts"))
            self.assertEqual(fonts.get(qn("w:ascii")), "Cordia New", text)
            self.assertEqual(fonts.get(qn("w:cs")), "Cordia New", text)
            self.assertEqual(properties.find(qn("w:sz")).get(qn("w:val")), "32", text)
            self.assertEqual(properties.find(qn("w:szCs")).get(qn("w:val")), "32", text)


class ParagraphFormatting(unittest.TestCase):
    def test_thai_justification_is_the_default(self):
        document, _ = build_document([Block(text="ทดสอบ")])
        pPr = document.paragraphs[0]._p.find(qn("w:pPr"))
        self.assertEqual(pPr.find(qn("w:jc")).get(qn("w:val")), "thaiDistribute")

    def test_alignment_can_be_overridden(self):
        document, _ = build_document([Block(text="หัวเรื่อง", align="center")])
        pPr = document.paragraphs[0]._p.find(qn("w:pPr"))
        self.assertEqual(pPr.find(qn("w:jc")).get(qn("w:val")), "center")

    def test_paragraph_mark_carries_font_and_size(self):
        document, _ = build_document([Block(text="ทดสอบ", size_pt=14)])
        rPr = document.paragraphs[0]._p.find(qn("w:pPr")).find(qn("w:rPr"))
        self.assertEqual(rPr.find(qn("w:rFonts")).get(qn("w:cs")), "Cordia New")
        self.assertEqual(rPr.find(qn("w:szCs")).get(qn("w:val")), "28")

    def test_bold_is_written_only_when_asked_for(self):
        plain, _ = build_document([Block(text="ทดสอบ")])
        bold, _ = build_document([Block(text="ทดสอบ", bold=True)])
        self.assertIsNone(plain.paragraphs[0].runs[0]._element.rPr.find(qn("w:b")))
        marked = bold.paragraphs[0].runs[0]._element.rPr
        self.assertIsNotNone(marked.find(qn("w:b")))
        self.assertIsNotNone(marked.find(qn("w:bCs")))

    def test_properties_are_in_schema_order(self):
        """Word rejects w:rPr children out of order, and szCs has no typed
        accessor in python-docx, so the order is worth asserting."""
        document, _ = build_document([Block(text="ทดสอบ", bold=True)])
        tags = [child.tag.split("}")[1] for child in document.paragraphs[0].runs[0]._element.rPr]
        self.assertEqual(tags, ["rFonts", "b", "bCs", "sz", "szCs", "cs"])

    def test_indent_spacing_and_page_break(self):
        document, _ = build_document([
            Block(text="ทดสอบ", first_line_indent_cm=1.27, space_after_pt=12, page_break_before=True)
        ])
        paragraph_format = document.paragraphs[0].paragraph_format
        self.assertEqual(round(paragraph_format.first_line_indent.cm, 2), 1.27)
        self.assertEqual(paragraph_format.space_after.pt, 12)
        self.assertTrue(paragraph_format.page_break_before)


class Layout(unittest.TestCase):
    def test_a4_portrait_with_one_inch_margins(self):
        document, _ = build_document([Block(text="ทดสอบ")])
        section = document.sections[0]
        self.assertEqual(round(section.page_width.mm), 210)
        self.assertEqual(round(section.page_height.mm), 297)
        for margin in (section.top_margin, section.right_margin, section.bottom_margin, section.left_margin):
            self.assertEqual(round(margin.mm, 1), 25.4)


class Templates(unittest.TestCase):
    def make_template(self, directory: Path) -> Path:
        """A stand-in for the firm's .docx: content to be cleared, and a style
        that has to survive the clearing."""
        template = Document()
        template.add_paragraph("content that must not survive")
        template.add_paragraph("nor this")
        template.styles.add_style("FirmStyle", 1)
        path = directory / "template.docx"
        template.save(str(path))
        return path

    def test_template_body_is_cleared_but_styles_are_kept(self):
        with TemporaryDirectory() as tmp:
            path = self.make_template(Path(tmp))
            document, used = build_document([Block(text="เนื้อหาใหม่")], template=str(path))
            self.assertEqual(used, path)
            self.assertEqual([p.text for p in document.paragraphs], ["เนื้อหาใหม่"])
            self.assertIn("FirmStyle", [s.name for s in document.styles])

    def test_section_properties_survive(self):
        with TemporaryDirectory() as tmp:
            path = self.make_template(Path(tmp))
            document, _ = build_document([Block(text="เนื้อหาใหม่")], template=str(path))
            self.assertEqual(len(document.sections), 1)

    def test_a_missing_template_is_reported(self):
        with self.assertRaises(FileNotFoundError):
            build_document([Block(text="ทดสอบ")], template="/nowhere/ABC.docx")

    def test_no_template_still_builds(self):
        document, used = build_document([Block(text="ทดสอบ")])
        self.assertIsNone(used)
        self.assertEqual(len(document.paragraphs), 1)


class Naming(unittest.TestCase):
    def test_file_name_falls_back_to_the_title(self):
        self.assertEqual(file_name_of("พินัยกรรม", None), "พินัยกรรม")
        self.assertEqual(file_name_of(None, None), "Document")
        self.assertEqual(file_name_of("a", "b/c:d"), "b c d")

    def test_a_second_document_lands_beside_the_first(self):
        with TemporaryDirectory() as tmp:
            out = Path(tmp)
            document, _ = build_document([Block(text="ทดสอบ")])
            first = save_document(document, out, "พินัยกรรม")
            second = save_document(document, out, "พินัยกรรม")
            self.assertEqual(first.name, "พินัยกรรม.docx")
            self.assertEqual(second.name, "พินัยกรรม (2).docx")
            self.assertEqual(free_path(out, "พินัยกรรม").name, "พินัยกรรม (3).docx")


if __name__ == "__main__":
    unittest.main()
