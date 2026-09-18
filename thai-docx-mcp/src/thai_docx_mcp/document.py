"""Build a Thai legal .docx — the engine, with no MCP in it.

Ported from the python-docx prototype. Two things make a Thai page look the way
a Thai lawyer expects, and neither is what python-docx writes by default:

  * ``w:jc="thaiDistribute"`` — Thai justification, which fills a line by
    spreading the characters rather than the word gaps Thai does not have;
  * ``w:cs`` and ``w:szCs`` on every Thai run, because Word treats Thai as
    complex script and otherwise sizes and styles it from the Latin properties.

Arabic numerals (700, 1,000.50, 2569/2026) are split into runs of their own and
left without ``w:cs`` — which is what Word itself writes when a Thai paragraph
carries digits.

Loading the firm's own .docx as a template is what preserves its styles,
settings, theme, font table and any embedded fonts; only the body is cleared
and rewritten. Without one the document still carries Cordia New on every run,
but a reader with no Cordia New installed will see a substitute.
"""

from __future__ import annotations

import os
import re
from pathlib import Path
from typing import Literal

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt
from docx.text.paragraph import Paragraph
from pydantic import BaseModel, Field

FONT_NAME = "Cordia New"
FONT_SIZE = 16.0

TEMPLATE_ENV = "JURATOOLS_DOCX_TEMPLATE"

ALIGNMENTS = {
    "thaiDistribute": WD_ALIGN_PARAGRAPH.THAI_JUSTIFY,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}

Align = Literal["thaiDistribute", "justify", "left", "center", "right"]

# 700 · 1,000 · 1,000.50 · 2569/2026 — one run, no w:cs.
NUMBER_PATTERN = r"[0-9]+(?:[.,/][0-9]+)*"


class Block(BaseModel):
    """One paragraph of the body. The single input contract, shared by the
    engine and the MCP tool so the two can never drift apart on what is valid."""

    text: str = Field(description="The paragraph text. Thai and Arabic numerals may be mixed freely")
    bold: bool = Field(default=False, description="Bold the whole paragraph")
    align: Align = Field(
        default="thaiDistribute",
        description='Paragraph alignment. "thaiDistribute" is the Thai justification used in legal drafting; use "center" for a heading',
    )
    size_pt: float | None = Field(default=None, description="Font size in points for this paragraph. Defaults to the document size")
    first_line_indent_cm: float | None = Field(default=None, description="First-line indent in centimetres, e.g. 1.27 for the usual half-inch")
    space_after_pt: float | None = Field(default=None, description="Blank space after the paragraph, in points")
    page_break_before: bool = Field(default=False, description="Start this paragraph on a new page")


# ── Run and paragraph properties ────────────────────────────────────────

def set_run_fonts(properties, font_name: str) -> None:
    """Set the font for Latin, high ANSI and complex-script text alike."""
    fonts = properties.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), font_name)
    fonts.set(qn("w:hAnsi"), font_name)
    fonts.set(qn("w:cs"), font_name)


def set_font_size(properties, size: float) -> None:
    """Set the ordinary and the complex-script size. Word stores half-points.

    python-docx has a typed accessor for ``w:sz`` but not for ``w:szCs``, and
    Word rejects the children of ``w:rPr`` out of schema order — where szCs
    immediately follows sz, which is what makes ``addnext`` the right insert.
    """
    half_points = str(round(size * 2))
    sz = properties.get_or_add_sz()
    sz.set(qn("w:val"), half_points)

    szCs = properties.find(qn("w:szCs"))
    if szCs is None:
        szCs = OxmlElement("w:szCs")
        sz.addnext(szCs)
    szCs.set(qn("w:val"), half_points)


def set_bold(properties, bold: bool) -> None:
    """Add the bold markers only when bold is wanted.

    Leaving out the explicit ``<w:b w:val="0"/>`` keeps the XML close to what
    Word writes by hand, which matters when the output is diffed against a
    document a lawyer produced in Word.
    """
    if not bold:
        return
    properties.get_or_add_b()
    properties.get_or_add_bCs()


def set_paragraph_formatting(paragraph: Paragraph, block: Block, font_name: str, font_size: float) -> None:
    """Carry the block's formatting onto the paragraph and its paragraph mark.

    The mark matters: Word sizes the line's leading from it, so a 16pt body with
    an 11pt paragraph mark comes out with uneven line spacing.
    """
    paragraph.alignment = ALIGNMENTS[block.align]

    size = block.size_pt or font_size
    paragraph_format = paragraph.paragraph_format
    if block.first_line_indent_cm is not None:
        paragraph_format.first_line_indent = Cm(block.first_line_indent_cm)
    if block.space_after_pt is not None:
        paragraph_format.space_after = Pt(block.space_after_pt)
    if block.page_break_before:
        paragraph_format.page_break_before = True

    paragraph_properties = paragraph._p.get_or_add_pPr()
    run_properties = paragraph_properties.find(qn("w:rPr"))
    if run_properties is None:
        run_properties = OxmlElement("w:rPr")
        # w:rPr is the last child of w:pPr but one — only w:sectPr follows it.
        paragraph_properties.insert_element_before(run_properties, "w:sectPr")
    set_run_fonts(run_properties, font_name)
    set_font_size(run_properties, size)
    set_bold(run_properties, block.bold)


def add_formatted_run(
    paragraph: Paragraph,
    text: str,
    font_name: str,
    size: float,
    bold: bool = False,
    complex_script: bool = False,
):
    """Add one run. ``complex_script`` adds the ``w:cs`` marker Thai needs."""
    run = paragraph.add_run(text)
    properties = run._element.get_or_add_rPr()
    set_run_fonts(properties, font_name)
    set_font_size(properties, size)
    set_bold(properties, bold)
    if complex_script:
        properties.get_or_add_cs()
    return run


def add_thai_with_arabic_numbers(
    paragraph: Paragraph,
    text: str,
    font_name: str = FONT_NAME,
    size: float = FONT_SIZE,
    bold: bool = False,
) -> Paragraph:
    """Split the text so numerals sit in runs of their own, without ``w:cs``."""
    for part in re.split(f"({NUMBER_PATTERN})", text):
        if not part:
            continue
        is_number = re.fullmatch(NUMBER_PATTERN, part) is not None
        add_formatted_run(
            paragraph=paragraph,
            text=part,
            font_name=font_name,
            size=size,
            bold=bold,
            complex_script=not is_number,
        )
    return paragraph


# ── The document ────────────────────────────────────────────────────────

def remove_document_content(document: DocumentObject) -> None:
    """Empty the body but keep its final ``w:sectPr``.

    Everything the template is loaded for — styles, settings, theme, font table,
    embedded fonts, the section configuration — lives outside the body or in
    that one element, so clearing the rest costs nothing.
    """
    body = document._element.body
    for element in list(body):
        if element.tag != qn("w:sectPr"):
            body.remove(element)


def ensure_a4_layout(document: DocumentObject) -> None:
    """A4 portrait with one-inch margins — the Thai legal page."""
    for section in document.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(25.4)
        section.right_margin = Mm(25.4)
        section.bottom_margin = Mm(25.4)
        section.left_margin = Mm(25.4)
        section.header_distance = Mm(12.7)
        section.footer_distance = Mm(12.7)


def resolve_template(template: str | os.PathLike[str] | None = None) -> Path | None:
    """Find the template: the call's own, else ``JURATOOLS_DOCX_TEMPLATE``.

    Returns None when neither is set — the caller then builds on python-docx's
    blank default and is told so, rather than failing.
    """
    candidate = template or os.environ.get(TEMPLATE_ENV) or None
    if candidate is None:
        return None
    path = Path(candidate).expanduser()
    if not path.exists():
        raise FileNotFoundError(f"Template not found: {path.resolve()}")
    return path


def build_document(
    blocks: list[Block],
    *,
    template: str | os.PathLike[str] | None = None,
    font_name: str = FONT_NAME,
    font_size: float = FONT_SIZE,
    enforce_a4: bool = True,
) -> tuple[DocumentObject, Path | None]:
    """Build the document from the blocks. Returns it with the template used."""
    template_path = resolve_template(template)

    if template_path is None:
        document = Document()
    else:
        document = Document(str(template_path))
    remove_document_content(document)
    if enforce_a4:
        ensure_a4_layout(document)

    for block in blocks:
        paragraph = document.add_paragraph()
        set_paragraph_formatting(paragraph, block, font_name, font_size)
        add_thai_with_arabic_numbers(
            paragraph=paragraph,
            text=block.text,
            font_name=font_name,
            size=block.size_pt or font_size,
            bold=block.bold,
        )

    return document, template_path


# ── Saving ──────────────────────────────────────────────────────────────

def safe_file_name(name: str) -> str:
    """A file name Windows, macOS and Linux all accept."""
    cleaned = re.sub(r'[\\/:*?"<>|]', " ", str(name))
    return re.sub(r"\s+", " ", cleaned).strip()


def file_name_of(title: str | None, file_name: str | None) -> str:
    if file_name:
        return safe_file_name(file_name)[:80] or "Document"
    if title:
        return safe_file_name(title)[:80] or "Document"
    return "Document"


def free_path(directory: Path, base: str) -> Path:
    """A document already written may have been sent out, so a second call with
    the same name lands beside it rather than on top of it."""
    index = 0
    while True:
        path = directory / (f"{base}.docx" if index == 0 else f"{base} ({index + 1}).docx")
        if not path.exists():
            return path
        index += 1


def save_document(document: DocumentObject, out_dir: Path, base_name: str) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    path = free_path(out_dir, base_name)
    document.save(str(path))
    return path
